import io
from app.db import get_db
from app.models.compra import Compra
from app.models.producto import Producto
from app.models.venta import Venta
from app.utils.auth import verificar_rol
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])


@router.get("/dashboard")
def obtener_resumen_reportes(
    db: Session = Depends(get_db),
    current_user=Depends(verificar_rol(["ADMIN"])),
):
  total_ventas_monto = (
      db.query(func.coalesce(func.sum(Venta.total), 0)).scalar()
  )
  total_compras_monto = (
      db.query(func.coalesce(func.sum(Compra.total), 0)).scalar()
  )
  cantidad_ventas = db.query(Venta).count()
  cantidad_compras = db.query(Compra).count()

  productos_stock_bajo = (
      db.query(Producto)
      .filter(Producto.existencia <= 5, Producto.estado == True)
      .all()
  )
  ventas_recientes = (
      db.query(Venta).order_by(Venta.fecha.desc()).limit(10).all()
  )
  compras_recientes = (
      db.query(Compra).order_by(Compra.fecha.desc()).limit(10).all()
  )

  return {
      "kpis": {
          "totalVentasMonto": float(total_ventas_monto),
          "totalComprasMonto": float(total_compras_monto),
          "gananciaBruta": float(total_ventas_monto - total_compras_monto),
          "cantidadVentas": cantidad_ventas,
          "cantidadCompras": cantidad_compras,
      },
      "stockBajo": [
          {
              "id": p.id,
              "codigo": p.codigo,
              "nombre": p.nombre,
              "existencia": p.existencia,
          }
          for p in productos_stock_bajo
      ],
      "ventasRecientes": [
          {
              "id": v.id,
              "numero": v.numero,
              "fecha": v.fecha,
              "total": float(v.total),
              "estado": v.estado,
              "detallesCount": len(v.detalles),
          }
          for v in ventas_recientes
      ],
      "comprasRecientes": [
          {
              "id": c.id,
              "numero": c.numero,
              "fecha": c.fecha,
              "total": float(c.total),
              "estado": c.estado,
              "proveedor": c.proveedor.nombre if c.proveedor else "N/A",
          }
          for c in compras_recientes
      ],
  }


@router.get("/exportar-excel")
def exportar_excel_inventario(
    db: Session = Depends(get_db),
    current_user=Depends(verificar_rol(["ADMIN"])),
):
  # 1. Obtener toda la información gerencial
  total_ventas_monto = (
      db.query(func.coalesce(func.sum(Venta.total), 0)).scalar()
  )
  total_compras_monto = (
      db.query(func.coalesce(func.sum(Compra.total), 0)).scalar()
  )
  cantidad_ventas = db.query(Venta).count()
  cantidad_compras = db.query(Compra).count()

  ventas_recientes = (
      db.query(Venta).order_by(Venta.fecha.desc()).limit(15).all()
  )
  productos = db.query(Producto).all()

  wb = openpyxl.Workbook()

  # Estilos corporativos profesionales
  title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
  section_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
  header_fill = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )
  header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
  zebra_fill = PatternFill(
      start_color="F9FBFD", end_color="F9FBFD", fill_type="solid"
  )
  alert_fill = PatternFill(
      start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
  )
  thin_border = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )

  # ==========================================
  # HOJA 1: RESUMEN EJECUTIVO Y KPIS
  # ==========================================
  ws1 = wb.active
  ws1.title = "Resumen Ejecutivo"
  ws1.views.sheetView[0].showGridLines = True

  ws1["A1"] = "SUPERMERCADO - PANEL DE RESUMEN Y ESTADÍSTICAS"
  ws1["A1"].font = title_font
  ws1.append([])

  # Tabla de KPIs / Métricas principales
  ws1["A3"] = "MÉTRICA / INDICADOR"
  ws1["B3"] = "VALOR GENERAL"
  ws1["A3"].fill = header_fill
  ws1["A3"].font = header_font
  ws1["B3"].fill = header_fill
  ws1["B3"].font = header_font
  ws1["A3"].border = thin_border
  ws1["B3"].border = thin_border

  kpis_data = [
      ("Ventas Totales (Monto)", float(total_ventas_monto)),
      ("Inversión en Compras (Monto)", float(total_compras_monto)),
      (
          "Flujo Neto / Margen Bruto",
          float(total_ventas_monto - total_compras_monto),
      ),
      ("Total Tickets de Venta Registrados", cantidad_ventas),
      ("Total Órdenes de Compra Registradas", cantidad_compras),
  ]

  for idx, (label, val) in enumerate(kpis_data, start=4):
    ws1[f"A{idx}"] = label
    ws1[f"B{idx}"] = val
    ws1[f"A{idx}"].border = thin_border
    ws1[f"B{idx}"].border = thin_border
    ws1[f"A{idx}"].font = Font(name="Calibri", size=10, bold=True)
    ws1[f"B{idx}"].font = Font(name="Calibri", size=10)
    if "Monto" in label or "Margen" in label:
      ws1[f"B{idx}"].number_format = '"Q"#,##0.00'
      ws1[f"B{idx}"].alignment = Alignment(horizontal="right")
    else:
      ws1[f"B{idx}"].alignment = Alignment(horizontal="center")

  ws1.append([])
  ws1.append([])

  # Tabla inferior: Últimas Ventas
  row_v = 11
  ws1.cell(row=row_v, column=1, value="ÚLTIMAS VENTAS REGISTRADAS").font = (
      section_font
  )
  row_v += 1
  v_headers = ["Ticket", "Fecha", "Items", "Estado", "Total"]
  for c_idx, h in enumerate(v_headers, start=1):
    cell = ws1.cell(row=row_v, column=c_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center")

  row_v += 1
  for v in ventas_recientes:
    ws1.cell(row=row_v, column=1, value=v.numero).alignment = Alignment(
        horizontal="center"
    )
    ws1.cell(row=row_v, column=2, value=str(v.fecha)).alignment = Alignment(
        horizontal="center"
    )
    ws1.cell(row=row_v, column=3, value=len(v.detalles)).alignment = Alignment(
        horizontal="center"
    )
    ws1.cell(row=row_v, column=4, value=v.estado).alignment = Alignment(
        horizontal="center"
    )
    t_cell = ws1.cell(row=row_v, column=5, value=float(v.total))
    t_cell.number_format = '"Q"#,##0.00'
    t_cell.alignment = Alignment(horizontal="right")

    for c_idx in range(1, 6):
      ws1.cell(row=row_v, column=c_idx).border = thin_border
      ws1.cell(row=row_v, column=c_idx).font = Font(name="Calibri", size=9)
    row_v += 1

  # ==========================================
  # HOJA 2: INVENTARIO DE PRODUCTOS
  # ==========================================
  ws2 = wb.create_sheet(title="Inventario de Productos")
  ws2.views.sheetView[0].showGridLines = True

  ws2["A1"] = "SUPERMERCADO - REPORTE GENERAL DE INVENTARIO"
  ws2["A1"].font = title_font
  ws2.append([])

  p_headers = [
      "ID",
      "Código",
      "Nombre del Producto",
      "Precio Compra",
      "Precio Venta",
      "Existencia",
      "Stock Mínimo",
      "Estado",
  ]
  ws2.append(p_headers)

  for col_num in range(1, len(p_headers) + 1):
    cell = ws2.cell(row=3, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = thin_border

  for row_idx, p in enumerate(productos, start=4):
    row_data = [
        p.id,
        p.codigo,
        p.nombre,
        float(p.precio_compra) if p.precio_compra else 0.0,
        float(p.precio_venta) if p.precio_venta else 0.0,
        p.existencia,
        p.stock_minimo,
        "Activo" if p.estado else "Inactivo",
    ]
    ws2.append(row_data)

    for col_idx in range(1, len(row_data) + 1):
      cell = ws2.cell(row=row_idx, column=col_idx)
      cell.border = thin_border
      cell.font = Font(name="Calibri", size=10)

      # Alerta visual en stock crítico
      if p.existencia <= 5:
        cell.fill = alert_fill

      if col_idx in [1, 6, 7]:
        cell.alignment = Alignment(horizontal="center", vertical="center")
      elif col_idx in [4, 5]:
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.number_format = '"Q"#,##0.00'
      else:
        cell.alignment = Alignment(horizontal="left", vertical="center")

  # Ajustar automáticamente el ancho de columnas en ambas pestañas
  for ws in [ws1, ws2]:
    for col in ws.columns:
      max_len = max(len(str(cell.value or "")) for cell in col)
      col_letter = get_column_letter(col[0].column)
      ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

  excel_file = io.BytesIO()
  wb.save(excel_file)
  excel_file.seek(0)

  return StreamingResponse(
      excel_file,
      media_type=(
          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
      ),
      headers={
          "Content-Disposition": (
              "attachment; filename=reporte_gerencial_supermercado.xlsx"
          )
      },
  )